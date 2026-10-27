"""Import a self-contained single-slide drawing as an editable PowerPoint group."""

from __future__ import annotations

import base64
from copy import deepcopy
from io import BytesIO
from pathlib import PurePosixPath
from typing import Any
from zipfile import ZipFile

from lxml import etree

from posterlib.delivery.pptx_fonts import import_embedded_fonts

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS = {"a": A, "p": P}
MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def decode_component(value: str) -> bytes:
    """Read the bounded native package embedded in a portable HTML component."""

    prefix = f"data:{MIME};base64,"
    if not value.startswith(prefix) or len(value) > 12 * 1024 * 1024:
        raise ValueError("Invalid native PPTX component data URI")
    return base64.b64decode(value[len(prefix):], validate=True)


def load_component(content: bytes) -> Any:
    """Read a bounded, offline, single-slide component without extracting its ZIP."""

    from pptx import Presentation

    with ZipFile(BytesIO(content)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 64 * 1024 * 1024:
            raise ValueError("PPTX component exceeds the expanded 64 MiB limit")
    presentation = Presentation(BytesIO(content))
    if len(presentation.slides) != 1:
        raise ValueError("PPTX component must contain exactly one slide")
    for part in presentation.part.package.iter_parts():
        if "vba" in part.content_type.lower():
            raise ValueError("PPTX component contains executable macro content")
        if any(rel.is_external for rel in part.rels.values()):
            raise ValueError("PPTX component must contain its resources locally")
    slide = presentation.slides[0]
    if slide.element.find(f"{{{P}}}timing") is not None:
        raise ValueError("PPTX component contains animation timing")
    if slide.element.xpath(".//p:ph"):
        raise ValueError("PPTX drawing components require explicit shapes instead of layout placeholders")
    if slide.element.xpath(".//a:tbl/a:tblPr/a:tableStyleId[normalize-space(.) != '']"):
        raise ValueError("PPTX component tables require explicit cell styling without an implicit table style")
    for parent in (slide.slide_layout, slide.slide_layout.slide_master):
        if any(not shape.is_placeholder for shape in parent.shapes):
            raise ValueError("PPTX drawing component has inherited layout artwork; place it on the slide")
    return presentation


def import_component(
    slide: Any, content: bytes, *, box: tuple[int, int, int, int], name: str,
) -> Any:
    """Fit the native drawing uniformly into ``box`` (EMU), retaining its objects."""

    from pptx.opc.constants import CONTENT_TYPE as CT
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.opc.package import Part, PartFactory
    from pptx.opc.packuri import PackURI
    from pptx.oxml import parse_xml

    source = load_component(content)
    source_slide = source.slides[0]
    theme = _source_theme(source)
    width, height = source.slide_width, source.slide_height
    left, top, available_width, available_height = box
    scale = min(available_width / width, available_height / height)
    target_width, target_height = round(width * scale), round(height * scale)
    package = slide.part.package
    import_embedded_fonts(source.part, package.presentation_part)
    occupied = {str(part.partname) for part in package.iter_parts()}
    copied: dict[Any, Any] = {}

    def partname(value: str) -> Any:
        path = PurePosixPath(value)
        candidate = str(path)
        index = 1
        while candidate in occupied:
            candidate = str(path.with_name(f"component-{index}-{path.name}"))
            index += 1
        occupied.add(candidate)
        return PackURI(candidate)

    def bind_relationships(root: Any, origin: Any, destination: Any = slide.part) -> None:
        mapping = {}
        for node in root.iter():
            for attribute, value in node.attrib.items():
                if attribute.startswith(f"{{{R}}}") and value not in mapping:
                    rel = origin.rels[value]
                    mapping[value] = destination.relate_to(copy_part(rel.target_part), rel.reltype)
        _remap_relationships(root, mapping)

    def copy_part(part: Any) -> Any:
        if part in copied:
            return copied[part]
        result = PartFactory(partname(str(part.partname)), part.content_type, package, part.blob)
        copied[part] = result
        relationships = {
            rel.rId: result.relate_to(copy_part(rel.target_part), rel.reltype)
            for rel in part.rels.values()
        }
        if relationships:
            root = etree.fromstring(part.blob, etree.XMLParser(resolve_entities=False, no_network=True))
            _remap_relationships(root, relationships)
            if hasattr(result, "_element"):
                result._element = parse_xml(etree.tostring(root))
            else:
                result.blob = etree.tostring(root)
        if part.content_type == CT.DML_CHART:
            chart_theme = _source_theme(source, part)
            _materialize_theme(result._element, chart_theme,
                               lambda node, origin: bind_relationships(node, origin, result))
            _scale_drawing(result._element, scale)
            if not any(rel.reltype == RT.THEME_OVERRIDE for rel in result.rels.values()):
                override = etree.Element(f"{{{A}}}themeOverride", nsmap={"a": A})
                override.extend(deepcopy(child) for child in chart_theme["root"].find(f"{{{A}}}themeElements"))
                theme_copy = Part(partname("/ppt/theme/component-theme.xml"), CT.OFC_THEME_OVERRIDE,
                                  package, etree.tostring(override))
                for child in override:
                    bind_relationships(child, chart_theme["owners"][child.tag], theme_copy)
                theme_copy.blob = etree.tostring(override)
                result.relate_to(theme_copy, RT.THEME_OVERRIDE)
        return result

    group = slide.shapes.add_group_shape()
    group.name = name
    _add_background(group, source, theme, bind_relationships, scale)
    children = [deepcopy(shape.element) for shape in source_slide.shapes]
    next_id = max(int(node.get("id")) for node in slide.element.iter(f"{{{P}}}cNvPr")) + 1
    ids: dict[str, str] = {}
    for child in children:
        for node in child.iter(f"{{{P}}}cNvPr"):
            ids[node.get("id")] = str(next_id)
            node.set("id", str(next_id))
            next_id += 1
    for child in children:
        bind_relationships(child, source_slide.part)
        _materialize_theme(child, theme, bind_relationships)
        _scale_drawing(child, scale)
        for tag in (f"{{{A}}}stCxn", f"{{{A}}}endCxn"):
            for connection in child.iter(tag):
                connection.set("id", ids[connection.get("id")])
        group.shapes._spTree.insert_element_before(child, "p:extLst")

    transform = group.element.grpSpPr.get_or_add_xfrm()
    transform.get_or_add_off().set("x", str(left + (available_width - target_width) // 2))
    transform.off.set("y", str(top + (available_height - target_height) // 2))
    transform.get_or_add_ext().set("cx", str(target_width))
    transform.ext.set("cy", str(target_height))
    transform.get_or_add_chOff().set("x", "0")
    transform.chOff.set("y", "0")
    transform.get_or_add_chExt().set("cx", str(target_width))
    transform.chExt.set("cy", str(target_height))
    return group


def _remap_relationships(root: Any, mapping: dict[str, str]) -> None:
    for node in root.iter():
        for attribute, value in list(node.attrib.items()):
            if attribute.startswith(f"{{{R}}}") and value in mapping:
                node.set(attribute, mapping[value])


def _scale_drawing(root: Any, scale: float) -> None:
    """Apply one uniform physical scale, keeping group coordinate maps unchanged."""

    dimensional_attributes = {
        "off": ("x", "y"), "ext": ("cx", "cy"),
        "chOff": ("x", "y"), "chExt": ("cx", "cy"),
        "rPr": ("sz", "spc"), "defRPr": ("sz", "spc"), "endParaRPr": ("sz", "spc"),
        "bodyPr": ("lIns", "rIns", "tIns", "bIns"), "ln": ("w",),
        "pPr": ("marL", "marR", "indent"), "spcPts": ("val",), "tab": ("pos",),
        "tr": ("h",), "gridCol": ("w",), "tcPr": ("marL", "marR", "marT", "marB"),
        "outerShdw": ("blurRad", "dist"), "innerShdw": ("blurRad", "dist"),
        "glow": ("rad",), "softEdge": ("rad",),
    }
    for node in root.iter():
        name = etree.QName(node)
        if name.namespace != A:
            continue
        if name.localname == "bodyPr":
            for attribute, default in (("lIns", 91440), ("rIns", 91440), ("tIns", 45720), ("bIns", 45720)):
                if attribute not in node.attrib:
                    node.set(attribute, str(default))
        for attribute in dimensional_attributes.get(name.localname, ()):
            if attribute in node.attrib:
                node.set(attribute, str(round(int(node.get(attribute)) * scale)))


def _source_theme(presentation: Any, owner: Any = None) -> dict[str, Any]:
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    slide = presentation.slides[0]
    master = slide.slide_layout.slide_master
    part = master.part.part_related_by(RT.THEME)
    root = etree.fromstring(part.blob)
    elements = root.find(f"{{{A}}}themeElements")
    owners = {child.tag: part for child in elements}
    for current in (slide.slide_layout.part, slide.part, owner):
        if current is None:
            continue
        for relation in current.rels.values():
            if relation.reltype not in {RT.THEME, RT.THEME_OVERRIDE}:
                continue
            override = etree.fromstring(relation.target_part.blob)
            replacement = override.find(f"{{{A}}}themeElements")
            for child in replacement if replacement is not None else override:
                previous = elements.find(child.tag)
                if previous is not None:
                    elements.replace(previous, deepcopy(child))
                    owners[child.tag] = relation.target_part
    colors = {etree.QName(node).localname: node[0] for node in root.find("a:themeElements/a:clrScheme", NS)}
    mapping = dict(master.element.find("p:clrMap", NS).attrib)
    for parent in (slide.slide_layout, slide):
        override = parent.element.find("p:clrMapOvr/a:overrideClrMapping", NS)
        if override is not None:
            mapping.update(override.attrib)
    colors.update({key: colors[value] for key, value in mapping.items()})
    return {
        "root": root, "colors": colors, "owners": owners,
        "text_style": master.element.find("p:txStyles/p:otherStyle", NS),
        "default_text_style": presentation.part._element.find(f"{{{P}}}defaultTextStyle"),
    }


_FILL_TAGS = ("noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill")
_PROPERTY_CHOICES = (
    _FILL_TAGS, ("custGeom", "prstGeom"), ("effectLst", "effectDag"),
    ("prstDash", "custDash"), ("round", "bevel", "miter"),
    ("uLnTx", "uLn"), ("uFillTx", "uFill"),
    ("buClrTx", "buClr"), ("buSzTx", "buSzPct", "buSzPts"),
    ("buFontTx", "buFont"), ("buNone", "buAutoNum", "buChar", "buBlip"),
)
_PROPERTY_ORDER = {
    "spPr": ("xfrm", "custGeom", "prstGeom", *_FILL_TAGS, "ln", "effectLst", "effectDag", "scene3d", "sp3d", "extLst"),
    "rPr": ("ln", *_FILL_TAGS, "effectLst", "effectDag", "highlight", "uLnTx", "uLn", "uFillTx", "uFill", "latin", "ea", "cs", "sym", "hlinkClick", "hlinkMouseOver", "rtl", "extLst"),
    "pPr": ("lnSpc", "spcBef", "spcAft", "buClrTx", "buClr", "buSzTx", "buSzPct", "buSzPts", "buFontTx", "buFont", "buNone", "buAutoNum", "buChar", "buBlip", "tabLst", "defRPr", "extLst"),
    "ln": (*_FILL_TAGS, "prstDash", "custDash", "round", "bevel", "miter", "headEnd", "tailEnd", "extLst"),
}


def _insert_property(parent: Any, value: Any) -> None:
    """Overlay a DrawingML property at its schema position, respecting fill choices."""

    name = etree.QName(value).localname
    choices = next((group for group in _PROPERTY_CHOICES if name in group), (name,))
    for child in list(parent):
        if etree.QName(child).localname in choices:
            parent.remove(child)
    local = etree.QName(parent).localname
    order = _PROPERTY_ORDER["rPr" if local.endswith("RPr") or local == "rPr" else local]
    position = order.index(name)
    for index, child in enumerate(parent):
        if order.index(etree.QName(child).localname) > position:
            parent.insert(index, value)
            return
    parent.append(value)


def _merge_properties(base: Any, override: Any, name: str = "pPr") -> Any:
    from pptx.oxml.xmlchemy import OxmlElement

    result = deepcopy(base) if base is not None else OxmlElement(f"a:{name}")
    result.tag = f"{{{A}}}{name}"
    if override is not None:
        result.attrib.update(override.attrib)
        for child in override:
            value = (_merge_properties(result.find(child.tag), child, "defRPr")
                     if child.tag == f"{{{A}}}defRPr" else deepcopy(child))
            _insert_property(result, value)
    return result


def _style_reference(ref: Any, theme: dict[str, Any], list_name: str, bind: Any) -> Any:
    index = int(ref.get("idx", "0"))
    if not index:
        return None
    if list_name == "fillStyleLst" and index >= 1001:
        list_name, index = "bgFillStyleLst", index - 1000
    values = theme["root"].find(f"a:themeElements/a:fmtScheme/a:{list_name}", NS)
    if values is None or not 0 < index <= len(values):
        raise ValueError("PPTX component has an unresolved theme style reference")
    value = deepcopy(values[index - 1])
    if list_name == "effectStyleLst":
        value = deepcopy(value.find(f"{{{A}}}effectLst"))
    for placeholder in list(value.iter(f"{{{A}}}schemeClr")):
        if placeholder.get("val") == "phClr" and len(ref):
            replacement = deepcopy(ref[0])
            replacement.extend(deepcopy(child) for child in placeholder)
            placeholder.getparent().replace(placeholder, replacement)
    bind(value, theme["owners"][f"{{{A}}}fmtScheme"])
    return value


def _materialize_theme(root: Any, theme: dict[str, Any], bind: Any) -> None:
    """Resolve style hierarchy before moving native objects between source parts."""

    for shape in root.iter():
        style = shape.find(f"{{{P}}}style")
        properties = shape.find(f"{{{P}}}spPr")
        if style is not None and properties is not None:
            for role, list_name, present in (
                ("fill", "fillStyleLst", _FILL_TAGS),
                ("ln", "lnStyleLst", ("ln",)),
                ("effect", "effectStyleLst", ("effectLst", "effectDag")),
            ):
                ref = style.find(f"{{{A}}}{role}Ref")
                local = next((properties.find(f"{{{A}}}{tag}") for tag in present
                              if properties.find(f"{{{A}}}{tag}") is not None), None)
                if ref is not None and (local is None or role == "ln"):
                    value = _style_reference(ref, theme, list_name, bind)
                    if value is not None:
                        if local is not None:
                            value = _merge_properties(value, local, "ln")
                        _insert_property(properties, value)
        for text in shape:
            if text.find(f"{{{A}}}bodyPr") is None:
                continue
            local_style = text.find(f"{{{A}}}lstStyle")
            for paragraph in text.findall(f"{{{A}}}p"):
                current = paragraph.find(f"{{{A}}}pPr")
                level = int(current.get("lvl", "0") if current is not None else 0) + 1
                merged = None
                for inherited in (theme["default_text_style"], theme["text_style"]):
                    if inherited is not None:
                        merged = _merge_properties(merged, inherited.find(f"{{{A}}}defPPr"))
                        merged = _merge_properties(merged, inherited.find(f"{{{A}}}lvl{level}pPr"))
                font_ref = style.find(f"{{{A}}}fontRef") if style is not None else None
                if font_ref is not None:
                    props = _merge_properties(merged.find(f"{{{A}}}defRPr"), None, "defRPr")
                    if len(font_ref):
                        fill = etree.Element(f"{{{A}}}solidFill")
                        fill.append(deepcopy(font_ref[0]))
                        _insert_property(props, fill)
                    family = {"major": "mj", "minor": "mn"}.get(font_ref.get("idx"))
                    if family:
                        for tag, suffix in (("latin", "lt"), ("ea", "ea"), ("cs", "cs")):
                            _insert_property(props, etree.Element(f"{{{A}}}{tag}", typeface=f"+{family}-{suffix}"))
                    _insert_property(merged, props)
                if local_style is not None:
                    merged = _merge_properties(merged, local_style.find(f"{{{A}}}defPPr"))
                    merged = _merge_properties(merged, local_style.find(f"{{{A}}}lvl{level}pPr"))
                merged = _merge_properties(merged, current)
                if current is not None:
                    paragraph.replace(current, merged)
                else:
                    paragraph.insert(0, merged)
        if style is not None:
            shape.remove(style)
    _resolve_theme_values(root, theme)


def _resolve_theme_values(root: Any, theme: dict[str, Any]) -> None:
    for node in list(root.iter(f"{{{A}}}schemeClr")):
        value = theme["colors"].get(node.get("val"))
        if value is None:
            raise ValueError(f"Unresolved component theme color: {node.get('val')}")
        replacement = deepcopy(value)
        if replacement.tag == f"{{{A}}}sysClr":
            replacement.tag = f"{{{A}}}srgbClr"
            replacement.set("val", replacement.attrib.pop("lastClr"))
        replacement.extend(deepcopy(child) for child in node)
        node.getparent().replace(node, replacement)
    for tag in ("latin", "ea", "cs"):
        for node in root.iter(f"{{{A}}}{tag}"):
            value = node.get("typeface", "")
            if value.startswith(("+mj-", "+mn-")):
                family = "majorFont" if value.startswith("+mj-") else "minorFont"
                source = theme["root"].find(f"a:themeElements/a:fontScheme/a:{family}/a:{tag}", NS)
                node.set("typeface", source.get("typeface", ""))


def _add_background(group: Any, source: Any, theme: dict[str, Any], bind: Any, scale: float) -> None:
    """Retain the source canvas as the group's bottom native object."""

    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE

    shape = group.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, source.slide_width, source.slide_height)
    shape.name = group.name + ".background"
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
    shape.line.fill.background()
    style = shape.element.find(f"{{{P}}}style")
    if style is not None:
        shape.element.remove(style)
    slide = source.slides[0]
    for owner in (slide, slide.slide_layout, slide.slide_layout.slide_master):
        background = owner.element.find(f"{{{P}}}cSld/{{{P}}}bg")
        if background is None:
            continue
        properties = background.find(f"{{{P}}}bgPr")
        if properties is not None:
            for child in properties:
                value = deepcopy(child)
                bind(value, owner.part)
                _insert_property(shape.element.spPr, value)
        else:
            value = _style_reference(background.find(f"{{{P}}}bgRef"), theme, "fillStyleLst", bind)
            if value is not None:
                _insert_property(shape.element.spPr, value)
        break
    _resolve_theme_values(shape.element, theme)
    _scale_drawing(shape.element, scale)
